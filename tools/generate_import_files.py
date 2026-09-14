import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

output_dir = r"c:\PowerApps\dados"
os.makedirs(output_dir, exist_ok=True)

# 1. Dados de Clientes
clientes_headers = ["Nome", "Email", "Telefone", "CPF", "Empresa", "Status", "DataNascimento", "Observacoes"]
clientes_data = [
    ["Ana Beatriz Souza", "ana.souza@techcorp.com.br", "(11) 98765-4321", "123.456.789-01", "TechCorp Solucoes", "Ativo", "1988-04-15", "Cliente VIP, plano Enterprise"],
    ["Carlos Eduardo Lima", "carlos.lima@inovare.com.br", "(21) 97654-3210", "234.567.890-12", "Inovare Logistica", "Ativo", "1992-08-22", "Contato preferencial por e-mail"],
    ["Mariana Costa Ribeiro", "mariana.costa@globalmed.com", "(31) 96543-2109", "345.678.901-23", "GlobalMed Saude", "Ativo", "1985-11-30", "Foco em suporte a sistemas criticos"],
    ["Fernando Albuquerque", "fernando@albuquerqueadv.com", "(41) 95432-1098", "456.789.012-34", "Albuquerque Advocacia", "Prospect", "1979-02-18", "Interesse em migracao de plano"],
    ["Juliana Martins Rocha", "juliana.martins@varejomax.com", "(51) 94321-0987", "567.890.123-45", "VarejoMax Brasil", "Ativo", "1995-07-09", "Abertura frequente de chamados operacionais"],
    ["Roberto Silveira Santos", "roberto@silveiraeng.com.br", "(61) 93210-9876", "678.901.234-56", "Silveira Engenharia", "Inativo", "1975-12-03", "Contrato suspenso temporariamente"],
    ["Patricia Mendes Duarte", "patricia@fintechpay.com", "(11) 92109-8765", "789.012.345-67", "FintechPay Brasil", "Ativo", "1990-09-25", "Cliente estrategico de alta prioridade"]
]

# 2. Dados de Atendimentos
atendimentos_headers = ["Titulo", "Protocolo", "Cliente", "Prioridade", "Status", "Canal", "DataAbertura", "Descricao"]
atendimentos_data = [
    ["Erro de sincronizacao no portal", "ATD-2026-0001", "Ana Beatriz Souza", "Alta", "Em andamento", "Portal", "2026-09-10 09:30", "Usuario relata falha ao sincronizar cadastros com a base principal."],
    ["Solicitacao de novo usuario de acesso", "ATD-2026-0002", "Carlos Eduardo Lima", "Baixa", "Resolvido", "E-mail", "2026-09-10 10:15", "Criacao de usuario de leitura para a equipe financeira."],
    ["Lentidao intermitente no fechamento", "ATD-2026-0003", "Mariana Costa Ribeiro", "Critica", "Aberto", "Telefone", "2026-09-11 14:00", "Relato de tempo de resposta acima de 15 segundos nas operacoes."],
    ["Duvida sobre faturamento do plano", "ATD-2026-0004", "Fernando Albuquerque", "Media", "Aguardando cliente", "Chat", "2026-09-12 11:20", "Solicitado espelho detalhado da fatura de agosto."],
    ["Falha na exportacao de relatorios PDF", "ATD-2026-0005", "Juliana Martins Rocha", "Media", "Em andamento", "Portal", "2026-09-13 16:45", "Botao de download nao gera o PDF com dados consolidados."],
    ["Revisao de permissoes de seguranca", "ATD-2026-0006", "Patricia Mendes Duarte", "Alta", "Aberto", "E-mail", "2026-09-14 08:30", "Necessidade de limitar visualizacao de dados sensiveis."]
]

# 3. Dados de Tarefas
tarefas_headers = ["Assunto", "Atendimento", "Status", "Previsao", "Descricao"]
tarefas_data = [
    ["Analisar logs de integracao no servidor", "ATD-2026-0001", "Em andamento", "2026-09-15", "Coletar telemetria e identificar timeout da API."],
    ["Enviar credenciais provisorias", "ATD-2026-0002", "Concluida", "2026-09-10", "Acesso gerado e enviado para o e-mail cadastrado."],
    ["Testar capacidade do banco e indices", "ATD-2026-0003", "Pendente", "2026-09-16", "Identificar queries de fechamento com alto consumo de CPU."],
    ["Validar modelo de geracao de PDF", "ATD-2026-0005", "Em andamento", "2026-09-17", "Verificar biblioteca de renderizacao de relatorios."],
    ["Mapear papeis e perfis da equipe", "ATD-2026-0006", "Pendente", "2026-09-18", "Definir matriz de permissoes com o cliente."]
]

# Gerar arquivos CSV (com codificação UTF-8 com BOM para Excel no Brasil reconhecer perfeitamente)
def save_csv(filename, headers, rows):
    path = os.path.join(output_dir, filename)
    with open(path, mode="w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(headers)
        writer.writerows(rows)
    print(f"CSV criado: {path}")

save_csv("clientes.csv", clientes_headers, clientes_data)
save_csv("atendimentos.csv", atendimentos_headers, atendimentos_data)
save_csv("tarefas.csv", tarefas_headers, tarefas_data)

# Gerar arquivo Excel (.xlsx) profissional com abas estilizadas
wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove aba padrao vazia

def add_styled_sheet(wb, title, headers, rows, header_color="1F4E79"):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True
    
    # Headers
    ws.append(headers)
    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    for col_num, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Rows
    row_font = Font(name="Segoe UI", size=10)
    for r_idx, row in enumerate(rows, 2):
        ws.append(row)
        for cell in ws[r_idx]:
            cell.font = row_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
            
    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

add_styled_sheet(wb, "Clientes", clientes_headers, clientes_data, "1F4E79")
add_styled_sheet(wb, "Atendimentos", atendimentos_headers, atendimentos_data, "2F5597")
add_styled_sheet(wb, "Tarefas", tarefas_headers, tarefas_data, "333f48")

xlsx_path = os.path.join(output_dir, "CustomerServiceData.xlsx")
wb.save(xlsx_path)
print(f"Excel criado: {xlsx_path}")
